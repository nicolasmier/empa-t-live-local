from fastapi import APIRouter, Depends, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.database import get_db
from app.errores import ErrorAPI
from app.models import Corredor, Distancia, Inscripcion, Usuario
from app.schemas.importacion import ImportacionFila, ImportacionResultado
from app.security.dependencies import requiere_admin
from app.servicios.auditoria import registrar_auditoria

router = APIRouter(tags=["importacion"])

# Columnas obligatorias del Excel (decisión clarify 2026-07-16, spec v10 Anexo B §4);
# "club" es opcional. La distancia se toma de la ruta, no de una columna.
# email/telefono son alternativos (decisión 2026-07-19): la cabecera debe traer
# al menos una de las dos, y cada fila debe rellenar al menos una — un email mal
# escrito en el Excel dejaba la inscripción huérfana sin forma de recuperarla; con
# las dos vías, si una falla la otra sigue sirviendo para el matching.
_COLUMNAS_BASE_REQUERIDAS = ("nombre", "apellidos", "dorsal")
_COLUMNAS_IDENTIDAD = ("email", "telefono")


def _valor_texto(fila: tuple, indice: dict[str, int], columna: str) -> str | None:
    i = indice.get(columna)
    if i is None or i >= len(fila) or fila[i] is None:
        return None
    valor = str(fila[i]).strip()
    return valor or None


def _valor_dorsal(fila: tuple, indice: dict[str, int]) -> str | None:
    # openpyxl devuelve las celdas numéricas como int/float, no str (p. ej. dorsal 101 -> 101.0).
    i = indice.get("dorsal")
    if i is None or i >= len(fila) or fila[i] is None:
        return None
    valor = fila[i]
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip() or None


def _leer_filas(archivo: UploadFile) -> tuple[dict[str, int], list[tuple]]:
    libro = load_workbook(filename=archivo.file, read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    filas_iter = hoja.iter_rows(values_only=True)
    try:
        cabecera = [str(c).strip().lower() if c is not None else "" for c in next(filas_iter)]
    except StopIteration as exc:
        raise ErrorAPI(400, "excel_vacio", "El fichero Excel no contiene filas.") from exc

    faltantes = [c for c in _COLUMNAS_BASE_REQUERIDAS if c not in cabecera]
    if faltantes:
        raise ErrorAPI(
            400,
            "columnas_faltantes",
            f"Faltan columnas obligatorias en el Excel: {', '.join(faltantes)}.",
        )
    if not any(c in cabecera for c in _COLUMNAS_IDENTIDAD):
        raise ErrorAPI(
            400,
            "columnas_faltantes",
            "Falta la columna 'email' o 'telefono' (al menos una es obligatoria).",
        )

    indice = {nombre: i for i, nombre in enumerate(cabecera)}
    return indice, list(filas_iter)


@router.post(
    "/distancias/{distancia_id}/corredores/importar",
    response_model=ImportacionResultado,
    status_code=201,
)
def importar_corredores(
    distancia_id: int,
    archivo: UploadFile,
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(requiere_admin),
) -> ImportacionResultado:
    distancia = db.get(Distancia, distancia_id)
    if distancia is None:
        raise ErrorAPI(404, "no_encontrado", "Distancia no encontrada.")

    try:
        indice, filas = _leer_filas(archivo)
    except (KeyError, ValueError) as exc:
        raise ErrorAPI(400, "excel_invalido", "El fichero no es un Excel válido.") from exc

    resultado_filas: list[ImportacionFila] = []
    contadores = {"creado": 0, "actualizado": 0, "sin_cambios": 0, "error": 0}

    for n, fila in enumerate(filas, start=2):  # la fila 1 es la cabecera
        if fila is None or all(c is None for c in fila):
            continue  # fila totalmente vacía: se ignora, no cuenta como error

        email = _valor_texto(fila, indice, "email")
        telefono = _valor_texto(fila, indice, "telefono")
        nombre = _valor_texto(fila, indice, "nombre")
        apellidos = _valor_texto(fila, indice, "apellidos")
        club = _valor_texto(fila, indice, "club")
        dorsal = _valor_dorsal(fila, indice)

        if not nombre or not apellidos or not dorsal or not (email or telefono):
            contadores["error"] += 1
            resultado_filas.append(
                ImportacionFila(
                    fila=n,
                    email=email,
                    telefono=telefono,
                    dorsal=dorsal,
                    estado="error",
                    motivo=(
                        "Faltan campos obligatorios (nombre, apellidos, dorsal, "
                        "y al menos uno de email/telefono)."
                    ),
                )
            )
            continue

        # email/telefono alternativos (decisión 2026-07-19): se busca primero por
        # email y, si no hay coincidencia, por telefono — así una fila con el email
        # mal escrito pero el teléfono correcto (o viceversa) sigue enlazando con la
        # cuenta ya existente en vez de crear un duplicado huérfano.
        corredor_existente = None
        if email:
            corredor_existente = db.query(Corredor).filter_by(email=email).one_or_none()
        if corredor_existente is None and telefono:
            corredor_existente = db.query(Corredor).filter_by(telefono=telefono).one_or_none()

        # Comprobación de colisión de dorsal ANTES de crear nada (si la fila es
        # inválida no debe dejar un `Corredor` huérfano a medio insertar): el
        # dorsal choca de verdad solo si pertenece a un corredor DISTINTO del que
        # esta fila identifica (por email o por teléfono).
        colision = (
            db.query(Inscripcion)
            .filter(Inscripcion.distancia_id == distancia_id, Inscripcion.dorsal == dorsal)
            .one_or_none()
        )
        if colision is not None and (
            corredor_existente is None or colision.corredor_id != corredor_existente.id
        ):
            contadores["error"] += 1
            resultado_filas.append(
                ImportacionFila(
                    fila=n,
                    email=email,
                    telefono=telefono,
                    dorsal=dorsal,
                    estado="error",
                    motivo=(
                        f"El dorsal '{dorsal}' ya está asignado a otro corredor "
                        "en esta distancia."
                    ),
                )
            )
            continue

        corredor = corredor_existente
        if corredor is None:
            # Cuenta pre-importada (RF-002): sin contraseña hasta que el corredor
            # complete el registro en la app con este mismo email (o, si el Excel
            # solo traía teléfono, aportando el email al registrarse).
            corredor = Corredor(
                nombre=nombre, apellidos=apellidos, club=club, email=email, telefono=telefono
            )
            db.add(corredor)
            db.flush()
        else:
            # No se toca nombre/apellidos/club de una cuenta ya existente (regla ya
            # establecida), pero sí se rellena email/telefono si la cuenta no lo
            # tenía todavía — cierra el hueco sin pisar nada que ya hubiera.
            if corredor.email is None and email:
                corredor.email = email
            if corredor.telefono is None and telefono:
                corredor.telefono = telefono

        inscripcion = (
            db.query(Inscripcion)
            .filter_by(corredor_id=corredor.id, distancia_id=distancia_id)
            .one_or_none()
        )
        if inscripcion is None:
            db.add(Inscripcion(corredor_id=corredor.id, distancia_id=distancia_id, dorsal=dorsal))
            db.flush()
            contadores["creado"] += 1
            estado_fila = "creado"
        elif inscripcion.dorsal != dorsal:
            inscripcion.dorsal = dorsal
            db.flush()
            contadores["actualizado"] += 1
            estado_fila = "actualizado"
        else:
            contadores["sin_cambios"] += 1
            estado_fila = "sin_cambios"

        resultado_filas.append(
            ImportacionFila(
                fila=n, email=email, telefono=telefono, dorsal=dorsal, estado=estado_fila
            )
        )

    registrar_auditoria(
        db,
        usuario.id,
        "corredores_importados",
        "distancia",
        str(distancia_id),
        detalle=dict(contadores),
    )
    db.commit()

    return ImportacionResultado(
        total_filas=len(resultado_filas),
        creados=contadores["creado"],
        actualizados=contadores["actualizado"],
        sin_cambios=contadores["sin_cambios"],
        errores=contadores["error"],
        filas=resultado_filas,
    )
